import unittest
from xml.etree import ElementTree
from zipfile import ZipFile
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

from backend.schemas import (
    DeliveryRunSheet, DeliveryRunSheetOrderSnapshot, DeliveryRunSheetTrip, ProductDetailLine,
)
from backend.services.delivery_run_sheet_excel_export_service import (
    DAILY_RUN_SHEET_HEADERS, build_delivery_run_sheet_excel,
    build_delivery_run_sheets_excel, delivery_run_sheet_product_display,
)


DATE = "2026-10-13"


def synthetic_order(number=1, lines=None, **changes):
    return replace(DeliveryRunSheetOrderSnapshot(
        row_id=f"QA-ROW-{number}", trip_no="trip1", row_no=number,
        task_type="ORDER", task_id=f"QA-ORDER-{number}",
        order_id_snapshot=f"QA-ORDER-{number}", invoice_number_snapshot=f"QA-INV-{number}",
        order_no_snapshot=None, company_name_snapshot=f"Synthetic Customer {number}",
        suburb_snapshot="QA Suburb", delivery_address_snapshot="Synthetic address",
        product_snapshot="Legacy snapshot product", pallet_quantity_snapshot=2,
        loose_bags_quantity_snapshot=0, note_snapshot=None,
        product_lines_snapshot=lines if lines is not None else [
            ProductDetailLine("COLOR TSHIRT RAGS", 450, "KG", "RSING", 45, "BAG10"),
        ],
    ), **changes)


def synthetic_sheet(orders=None, trip="trip1", **changes):
    orders = orders if orders is not None else [synthetic_order()]
    return replace(DeliveryRunSheet(
        run_sheet_id=f"QA-SHEET-{trip}", dispatch_date=DATE, delivery_date=DATE,
        driver_id="QA-DRIVER", driver_name_snapshot="Synthetic Driver",
        vehicle_id="QA-VEHICLE", vehicle_rego_snapshot="QA-ONLY",
        total_pallets=sum(order.pallet_quantity_snapshot for order in orders),
        total_loose_bags=0, status="SAVED", generated_at="2026-10-13T00:00:00Z",
        saved_at="2026-10-13T00:01:00Z", saved_by_account_name="Synthetic QA",
        saved_by_account_id=1, legacy_summary_id=None,
        trips=[DeliveryRunSheetTrip(trip or "trip1", orders)], trip_no=trip,
    ), **changes)


class DeliveryRunSheetExcelPaginationTest(unittest.TestCase):
    def workbook(self, sheet, date=False):
        content = (build_delivery_run_sheets_excel([sheet], DATE) if date
                   else build_delivery_run_sheet_excel(sheet))
        workbook = load_workbook(BytesIO(content))
        self.addCleanup(workbook.close)
        return workbook

    def test_snapshot_name_precedes_code_and_legacy(self):
        self.assertEqual("COLOR TSHIRT RAGS - 10kg x 45Bag",
                         delivery_run_sheet_product_display(synthetic_order()))

    def test_missing_or_whitespace_name_falls_back_to_recorded_code(self):
        for name in (None, "", " \t "):
            with self.subTest(name=name):
                order = synthetic_order(lines=[ProductDetailLine(name, 450, "KG", " RSING ", 45, "BAG10")])
                self.assertEqual("RSING - 10kg x 45Bag", delivery_run_sheet_product_display(order))

    def test_legacy_fallback_when_no_valid_line_identity(self):
        for lines in ([], [ProductDetailLine(" ", 1, "EA", None)]):
            with self.subTest(lines=lines):
                self.assertEqual("Legacy snapshot product",
                                 delivery_run_sheet_product_display(synthetic_order(lines=lines)))

    def test_explicit_bag_weights_and_quantities_have_no_trailing_zeros(self):
        for unit, quantity, expected in (
            ("BAG10", 45, "10kg x 45Bag"), ("BAG10", 90, "10kg x 90Bag"),
            ("BAG5", 1, "5kg x 1Bag"), ("BAG1.5", 12, "1.5kg x 12Bag"),
            (" bag1.50 ", 12.0, "1.5kg x 12Bag"), ("BaG 5.00", 1, "5kg x 1Bag"),
            ("BAG10", 0, "10kg x 0Bag"), ("BAG5", 1.5, "5kg x 1.5Bag"),
        ):
            with self.subTest(unit=unit, quantity=quantity):
                order = synthetic_order(lines=[ProductDetailLine("Snapshot name", 1, "EA", "CODE", quantity, unit)])
                self.assertEqual(f"Snapshot name - {expected}", delivery_run_sheet_product_display(order))

    def test_other_units_keep_meaningful_quantity_without_guessed_weights(self):
        for unit in ("PKT25", "CTN", "BAG", "BAGS", "EA", "EACH", "PALLETS",
                     "BAG0", "BAG-10", "BAG10OLD", "BAG10KG"):
            with self.subTest(unit=unit):
                order = synthetic_order(lines=[ProductDetailLine("Snapshot name", 100, "EA", "CODE", 4, unit)])
                self.assertEqual(f"Snapshot name - 4 {unit}", delivery_run_sheet_product_display(order))

    def test_missing_invalid_or_nonfinite_package_quantity_uses_base_quantity(self):
        for quantity in (None, "invalid", float("nan"), float("inf"), True):
            with self.subTest(quantity=quantity):
                order = synthetic_order(lines=[ProductDetailLine("Snapshot name", 100, "KG", "CODE", quantity, "BAG10")])
                self.assertEqual("Snapshot name - 100 KG", delivery_run_sheet_product_display(order))

    def test_negative_package_quantity_is_not_converted_to_weight_notation(self):
        order = synthetic_order(lines=[ProductDetailLine("Snapshot name", 1, "EA", "CODE", -2, "BAG10")])
        self.assertEqual("Snapshot name - -2 BAG10", delivery_run_sheet_product_display(order))

    def test_multiple_products_stay_in_one_cell_in_snapshot_order(self):
        lines = [ProductDetailLine("First snapshot product", 450, "KG", "ONE", 45, "BAG10"),
                 ProductDetailLine("Second snapshot product", 12, "EA", "TWO", 2, "CTN")]
        worksheet = self.workbook(synthetic_sheet([synthetic_order(lines=lines)])).active
        self.assertEqual("First snapshot product - 10kg x 45Bag\nSecond snapshot product - 2 CTN", worksheet["E9"].value)
        self.assertTrue(worksheet["E9"].alignment.wrap_text)

    def test_weight_is_independent_numeric_and_not_double_counted(self):
        lines = [ProductDetailLine("Direct recorded kilograms", 450, "KG", "ITEM10KG", 45, "BAG10"),
                 ProductDetailLine("Weight encoded in code", 18, "BAG", "ITEM5KG", 18, "BAG10")]
        worksheet = self.workbook(synthetic_sheet([synthetic_order(lines=lines)])).active
        self.assertEqual(540, worksheet["F9"].value)
        self.assertEqual(2, worksheet["G9"].value)
        self.assertIn("10kg x 18Bag", worksheet["E9"].value)

    def test_unknown_weight_remains_blank(self):
        order = synthetic_order(lines=[ProductDetailLine("Packet product", 250, "EA", "PACKET", 10, "PKT25")])
        worksheet = self.workbook(synthetic_sheet([order])).active
        self.assertIsNone(worksheet["F9"].value)
        self.assertEqual(2, worksheet["G9"].value)

    def test_short_rows_keep_handwriting_space_and_original_form_fields(self):
        worksheet = self.workbook(synthetic_sheet([synthetic_order(lines=[ProductDetailLine("Short", 1, "EA")])])).active
        self.assertEqual(21.95, worksheet.row_dimensions[9].height)
        self.assertEqual(21.95, worksheet.row_dimensions[26].height)
        self.assertEqual(DAILY_RUN_SHEET_HEADERS, [cell.value for cell in worksheet[8]])
        self.assertEqual("DRIVER: Synthetic Driver", worksheet["F1"].value)
        self.assertEqual("REGO #: QA-ONLY", worksheet["L1"].value)
        self.assertEqual("TRIP: TRIP 1", worksheet["F2"].value)
        self.assertEqual({"B5:E5", "F5:N5"}, {str(item) for item in worksheet.merged_cells.ranges})
        for column in range(8, 15):
            self.assertIsNone(worksheet.cell(9, column).value)

    def test_five_long_descriptions_have_room_for_the_final_line(self):
        lines = [ProductDetailLine(f"Synthetic product {number} with an extended descriptive snapshot name FINAL-{number}",
                                   10, "KG", f"CODE-{number}", 1, "BAG10") for number in range(1, 6)]
        worksheet = self.workbook(synthetic_sheet([synthetic_order(lines=lines)])).active
        self.assertEqual(5, len(worksheet["E9"].value.split("\n")))
        for line in lines:
            self.assertIn(line.product_name, worksheet["E9"].value)
        self.assertTrue(worksheet["E9"].value.endswith("FINAL-5 - 10kg x 1Bag"))
        self.assertGreater(worksheet.row_dimensions[9].height, 60)
        self.assertLessEqual(worksheet.row_dimensions[9].height, 409.5)

    def test_customer_suburb_invoice_and_product_wrapping_expand_rows(self):
        base = synthetic_order(lines=[ProductDetailLine("Short", 1, "EA")])
        ordinary = self.workbook(synthetic_sheet([base])).active.row_dimensions[9].height
        for field in ("company_name_snapshot", "suburb_snapshot", "invoice_number_snapshot"):
            with self.subTest(field=field):
                worksheet = self.workbook(synthetic_sheet([replace(base, **{field: "Synthetic long text " * 15})])).active
                self.assertGreater(worksheet.row_dimensions[9].height, ordinary)
        long_product = replace(base, product_lines_snapshot=[ProductDetailLine("Synthetic long product " * 15, 1, "EA")])
        self.assertGreater(self.workbook(synthetic_sheet([long_product])).active.row_dimensions[9].height, 60)

    def test_explicit_line_breaks_are_included_in_height(self):
        order = synthetic_order(lines=[ProductDetailLine("First line\nSecond line\nFinal line", 1, "EA")])
        worksheet = self.workbook(synthetic_sheet([order])).active
        self.assertIn("Final line", worksheet["E9"].value)
        self.assertGreaterEqual(worksheet.row_dimensions[9].height, 3 * 10.5 + 4)

    def test_wide_glyphs_require_more_space_than_narrow_glyphs(self):
        heights = []
        for name in ("i" * 80, "W" * 80, "产品" * 40):
            worksheet = self.workbook(synthetic_sheet([synthetic_order(lines=[ProductDetailLine(name, 1, "EA")])])).active
            heights.append(worksheet.row_dimensions[9].height)
        self.assertGreater(heights[1], heights[0])
        self.assertGreater(heights[2], heights[0])

    def test_18_30_and_55_customers_preserve_every_order_and_print_scope(self):
        for count in (18, 30, 55):
            with self.subTest(count=count):
                worksheet = self.workbook(synthetic_sheet([synthetic_order(number) for number in range(1, count + 1)])).active
                self.assertEqual(list(range(1, count + 1)), [worksheet.cell(row, 1).value for row in range(9, count + 9)])
                self.assertEqual([f"QA-INV-{number}" for number in range(1, count + 1)],
                                 [worksheet.cell(row, 4).value for row in range(9, count + 9)])
                self.assertEqual(450 * count, sum(worksheet.cell(row, 6).value for row in range(9, count + 9)))
                self.assertEqual(2 * count, sum(worksheet.cell(row, 7).value for row in range(9, count + 9)))
                self.assertEqual(0, worksheet.page_setup.fitToHeight)
                self.assertEqual(1, worksheet.page_setup.fitToWidth)
                self.assertEqual("landscape", worksheet.page_setup.orientation)
                self.assertEqual(str(worksheet.PAPERSIZE_A4), str(worksheet.page_setup.paperSize))
                self.assertEqual("$7:$8", worksheet.print_title_rows)
                self.assertEqual(14, worksheet.max_column)
                self.assertIn(f"$A$1:$N${count + 10}", worksheet.print_area)
                self.assertEqual("FINISH TIME:____________________________________", worksheet.cell(count + 10, 2).value)
                self.assertEqual(count + 10, worksheet.max_row)
                self.assertEqual([], worksheet.row_breaks.brk)

    def test_over_page_product_uses_labelled_continuations_without_repeating_totals(self):
        name = "\n".join(f"Synthetic extended product line {number} FINAL-{number}" for number in range(1, 121))
        order = synthetic_order(lines=[ProductDetailLine(name, 10, "KG")])
        sheet = synthetic_sheet([order])
        before = deepcopy(sheet)
        worksheet = self.workbook(sheet).active
        continuation_rows = [row for row in range(10, worksheet.max_row) if isinstance(worksheet.cell(row, 1).value, str)]
        self.assertGreater(len(continuation_rows), 0)
        rows = [9] + continuation_rows
        content = "".join(str(worksheet.cell(row, 5).value or "") for row in rows)
        self.assertEqual(delivery_run_sheet_product_display(order), content)
        self.assertIn("FINAL-120", content)
        self.assertEqual(10, sum(worksheet.cell(row, 6).value or 0 for row in rows))
        self.assertEqual(2, sum(worksheet.cell(row, 7).value or 0 for row in rows))
        for row in continuation_rows:
            self.assertEqual("1\ncont.", worksheet.cell(row, 1).value)
            self.assertIsNone(worksheet.cell(row, 6).value)
            self.assertIsNone(worksheet.cell(row, 7).value)
        self.assertTrue(all(worksheet.row_dimensions[row].height <= 409.5 for row in rows))
        self.assertEqual(before, sheet)

    def test_over_page_customer_text_is_preserved_in_continuations(self):
        customer = "Synthetic extraordinarily long customer name " * 100
        worksheet = self.workbook(synthetic_sheet([synthetic_order(company_name_snapshot=customer)])).active
        rows = [9] + [row for row in range(10, worksheet.max_row) if isinstance(worksheet.cell(row, 1).value, str)]
        text = "".join(str(worksheet.cell(row, 2).value or "") for row in rows)
        self.assertEqual(customer, text)
        self.assertGreater(len(rows), 1)

    def test_over_page_unbroken_product_text_is_not_truncated_or_rewritten(self):
        order = synthetic_order(lines=[ProductDetailLine("W" * 2000 + "FINAL", 1, "EA")])
        worksheet = self.workbook(synthetic_sheet([order])).active
        rows = [9] + [row for row in range(10, worksheet.max_row) if isinstance(worksheet.cell(row, 1).value, str)]
        content = "".join(str(worksheet.cell(row, 5).value or "") for row in rows)
        self.assertEqual(delivery_run_sheet_product_display(order), content)
        self.assertGreater(len(rows), 1)

    def test_multiline_order_uses_only_handwriting_rows_that_fit(self):
        order = synthetic_order(lines=[ProductDetailLine("Synthetic extended product description " * 10, 1, "EA")])
        worksheet = self.workbook(synthetic_sheet([order])).active
        self.assertLess(worksheet.max_row, 28)
        self.assertEqual(delivery_run_sheet_product_display(order), worksheet["E9"].value)
        self.assertEqual("FINISH TIME:____________________________________", worksheet.cell(worksheet.max_row, 2).value)
        self.assertEqual(0, worksheet.page_setup.fitToHeight)

    def test_over_page_continuations_are_not_followed_by_blank_customer_padding(self):
        order = synthetic_order(lines=[ProductDetailLine("Synthetic text\n" * 120, 10, "KG")])
        worksheet = self.workbook(synthetic_sheet([order])).active
        continuation_rows = [row for row in range(10, worksheet.max_row) if isinstance(worksheet.cell(row, 1).value, str)]
        self.assertGreater(len(continuation_rows), 0)
        self.assertEqual(continuation_rows[-1] + 2, worksheet.max_row)
        self.assertEqual([1], [worksheet.cell(row, 1).value for row in range(9, worksheet.max_row)
                              if isinstance(worksheet.cell(row, 1).value, int)])

    def test_single_date_trip1_trip2_and_legacy_exports_share_layout(self):
        for trip in ("trip1", "trip2", None):
            with self.subTest(trip=trip):
                sheet = synthetic_sheet(trip=trip)
                single = self.workbook(sheet).active
                date = self.workbook(sheet, date=True).active
                self.assertEqual(list(single.values), list(date.values))
                for worksheet in (single, date):
                    self.assertEqual(0, worksheet.page_setup.fitToHeight)
                    self.assertEqual(1, worksheet.page_setup.fitToWidth)
                    self.assertEqual("$7:$8", worksheet.print_title_rows)
                    self.assertEqual("COLOR TSHIRT RAGS - 10kg x 45Bag", worksheet["E9"].value)

    def test_legacy_combined_trips_keep_order_sequence(self):
        sheet = synthetic_sheet(trip=None)
        sheet.trips.append(DeliveryRunSheetTrip("trip2", [synthetic_order(2)]))
        worksheet = self.workbook(sheet).active
        self.assertIsNone(worksheet["F2"].value)
        self.assertEqual(["QA-INV-1", "QA-INV-2"], [worksheet.cell(row, 4).value for row in (9, 10)])
        self.assertEqual([1, 2], [worksheet.cell(row, 1).value for row in (9, 10)])

    def test_both_exports_leave_saved_snapshots_and_product_lines_unchanged(self):
        sheet = synthetic_sheet()
        before = deepcopy(sheet)
        self.workbook(sheet)
        self.workbook(sheet, date=True)
        self.assertEqual(before, sheet)
        self.assertEqual("COLOR TSHIRT RAGS", sheet.trips[0].orders[0].product_lines_snapshot[0].product_name)


class DeliveryRunSheetRedQuantityTest(unittest.TestCase):
    def workbook(self, sheet, date=False, rich_text=True):
        content = (build_delivery_run_sheets_excel([sheet], DATE) if date
                   else build_delivery_run_sheet_excel(sheet))
        workbook = load_workbook(BytesIO(content), rich_text=rich_text)
        self.addCleanup(workbook.close)
        return workbook

    def characters(self, cell):
        self.assertIsInstance(cell.value, CellRichText)
        colors = []
        for block in cell.value:
            if isinstance(block, TextBlock):
                self.assertTrue(block.font.b)
                self.assertEqual(7.5, block.font.sz)
                self.assertEqual("rgb", block.font.color.type)
                color = block.font.color.rgb[-6:]
                text = block.text
            else:
                color, text = "000000", block
            colors.extend([color] * len(text))
        self.assertTrue(cell.font.bold)
        self.assertEqual(7.5, cell.font.sz)
        self.assertTrue(cell.alignment.wrap_text)
        self.assertEqual("top", cell.alignment.vertical)
        self.assertFalse(cell.alignment.shrink_to_fit)
        self.assertTrue(cell.font.color is None or cell.font.color.rgb[-6:] != "C00000")
        return str(cell.value), colors

    def assert_quantity(self, order, name, quantity):
        cell = self.workbook(synthetic_sheet([order])).active["E9"]
        text, colors = self.characters(cell)
        prefix = name + " - "
        self.assertEqual(prefix + quantity, text)
        self.assertEqual(["000000"] * len(prefix) + ["C00000"] * len(quantity), colors)

    def test_one_product_colors_only_the_quantity_suffix(self):
        self.assert_quantity(synthetic_order(), "COLOR TSHIRT RAGS", "10kg x 45Bag")

    def test_bag10_bag5_and_bag_decimal_quantities_are_red(self):
        for unit, count, expected in (("BAG10", 45, "10kg x 45Bag"),
                                      ("BAG5", 1, "5kg x 1Bag"),
                                      ("BAG1.5", 12, "1.5kg x 12Bag")):
            with self.subTest(unit=unit):
                self.assert_quantity(synthetic_order(lines=[ProductDetailLine("Snapshot name", 200, "KG", "CODE", count, unit)]),
                                     "Snapshot name", expected)

    def test_packet_carton_bag_bags_and_each_quantities_are_red(self):
        for unit in ("PKT25", "CTN", "BAG", "BAGS", "EA", "EACH", "PALLETS"):
            with self.subTest(unit=unit):
                self.assert_quantity(synthetic_order(lines=[ProductDetailLine("Snapshot name", 200, "KG", "CODE", 4, unit)]),
                                     "Snapshot name", f"4 {unit}")

    def test_direct_kilogram_quantity_is_red(self):
        self.assert_quantity(synthetic_order(lines=[ProductDetailLine("PRODUCT NAME", 200, "KG")]),
                             "PRODUCT NAME", "200 KG")

    def test_numbers_and_separators_within_product_names_stay_black(self):
        name = "WORKSHOP MIX #29 - 10KG NET 40 X 40 - 200 KG"
        self.assert_quantity(synthetic_order(lines=[ProductDetailLine(name, 300, "KG", "CODE", 30, "BAG10")]),
                             name, "10kg x 30Bag")

    def test_missing_name_keeps_code_black_and_quantity_red(self):
        self.assert_quantity(synthetic_order(lines=[ProductDetailLine(" ", 200, "KG", "CODE29", 20, "BAG10")]),
                             "CODE29", "10kg x 20Bag")

    def test_multiple_product_lines_and_newline_colors_are_preserved(self):
        order = synthetic_order(lines=[ProductDetailLine("WORKSHOP MIX #29", 300, "KG", "CODE", 30, "BAG10"),
                                       ProductDetailLine("Product B", 2, "BAGS")])
        text, colors = self.characters(self.workbook(synthetic_sheet([order])).active["E9"])
        self.assertEqual("WORKSHOP MIX #29 - 10kg x 30Bag\nProduct B - 2 BAGS", text)
        expected = (["000000"] * len("WORKSHOP MIX #29 - ") + ["C00000"] * len("10kg x 30Bag")
                    + ["000000"] * len("\nProduct B - ") + ["C00000"] * len("2 BAGS"))
        self.assertEqual(expected, colors)

    def test_missing_quantity_and_legacy_text_remain_black_without_guessing(self):
        for order in (synthetic_order(lines=[ProductDetailLine("Product #29 10KG", None, None)]),
                      synthetic_order(lines=[ProductDetailLine("Product #29", None, "KG")]),
                      synthetic_order(lines=[], product_snapshot="Legacy product - 10kg x 45Bag")):
            with self.subTest(text=delivery_run_sheet_product_display(order)):
                text, colors = self.characters(self.workbook(synthetic_sheet([order])).active["E9"])
                self.assertEqual(delivery_run_sheet_product_display(order), text)
                self.assertEqual(["000000"] * len(text), colors)

    def test_multiple_product_xml_has_no_formatted_whitespace_only_runs(self):
        order = synthetic_order(lines=[ProductDetailLine("Product A #29", 300, "KG", "CODE", 30, "BAG10"),
                                       ProductDetailLine("Product B 10KG", 2, "BAGS")])
        content = build_delivery_run_sheet_excel(synthetic_sheet([order]))
        with ZipFile(BytesIO(content)) as archive:
            root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        cell = root.find(".//s:c[@r='E9']", namespace)
        for run in cell.findall("s:is/s:r", namespace):
            if run.find("s:rPr", namespace) is not None:
                self.assertTrue(run.find("s:t", namespace).text.strip())
        self.assertEqual(delivery_run_sheet_product_display(order),
                         self.workbook(synthetic_sheet([order]), rich_text=False).active["E9"].value)

    def test_long_multiline_name_continues_black_with_only_final_quantity_red(self):
        name = "\n".join(f"Synthetic product #29 10KG NET line {number}" for number in range(120))
        order = synthetic_order(lines=[ProductDetailLine(name, 10, "KG")])
        worksheet = self.workbook(synthetic_sheet([order])).active
        cells = [worksheet.cell(row, 5) for row in range(9, worksheet.max_row) if worksheet.cell(row, 5).value]
        self.assertGreater(len(cells), 1)
        text, colors = "", []
        for cell in cells:
            fragment, fragment_colors = self.characters(cell)
            text += fragment
            colors.extend(fragment_colors)
        self.assertEqual(delivery_run_sheet_product_display(order), text)
        self.assertEqual(["000000"] * len(name + " - ") + ["C00000"] * len("10 KG"), colors)
        for row in range(10, worksheet.max_row):
            if worksheet.cell(row, 1).value == "1\ncont.":
                self.assertTrue(worksheet.cell(row, 1).font.color is None)
        self.assertEqual(10, sum(worksheet.cell(row, 6).value or 0 for row in range(9, worksheet.max_row)))
        self.assertEqual(2, sum(worksheet.cell(row, 7).value or 0 for row in range(9, worksheet.max_row)))

    def test_quantity_run_can_span_multiple_continuations_without_text_loss(self):
        quantity = "9 UNIT-" + "X" * 2000
        order = synthetic_order(lines=[ProductDetailLine("Snapshot name", 10, "KG", "CODE", 9, "UNIT-" + "X" * 2000)])
        worksheet = self.workbook(synthetic_sheet([order])).active
        cells = [worksheet.cell(row, 5) for row in range(9, worksheet.max_row) if worksheet.cell(row, 5).value]
        self.assertGreater(len(cells), 1)
        text, colors = "", []
        for cell in cells:
            fragment, fragment_colors = self.characters(cell)
            text += fragment
            colors.extend(fragment_colors)
        self.assertEqual("Snapshot name - " + quantity, text)
        self.assertEqual(["000000"] * len("Snapshot name - ") + ["C00000"] * len(quantity), colors)

    def test_single_date_and_legacy_exports_keep_colors_and_plain_reader_content(self):
        for trip in ("trip1", "trip2", None):
            for date in (False, True):
                with self.subTest(trip=trip, date=date):
                    sheet = synthetic_sheet(trip=trip)
                    before = deepcopy(sheet)
                    rich = self.workbook(sheet, date=date).active
                    plain = self.workbook(sheet, date=date, rich_text=False).active
                    text, colors = self.characters(rich["E9"])
                    self.assertEqual("COLOR TSHIRT RAGS - 10kg x 45Bag", text)
                    self.assertEqual(text, plain["E9"].value)
                    self.assertIsInstance(plain["E9"].value, str)
                    self.assertIn("C00000", colors)
                    self.assertEqual((450, 2), (rich["F9"].value, rich["G9"].value))
                    self.assertEqual(0, rich.page_setup.fitToHeight)
                    self.assertEqual(1, rich.page_setup.fitToWidth)
                    self.assertEqual("$7:$8", rich.print_title_rows)
                    self.assertEqual(before, sheet)

    def test_18_30_55_orders_keep_plain_text_totals_and_page_setup(self):
        for count in (18, 30, 55):
            with self.subTest(count=count):
                worksheet = self.workbook(synthetic_sheet([synthetic_order(number) for number in range(1, count + 1)])).active
                for row in range(9, count + 9):
                    text, colors = self.characters(worksheet.cell(row, 5))
                    self.assertEqual("COLOR TSHIRT RAGS - 10kg x 45Bag", text)
                    self.assertEqual(["000000"] * len("COLOR TSHIRT RAGS - ") + ["C00000"] * len("10kg x 45Bag"), colors)
                self.assertEqual(450 * count, sum(worksheet.cell(row, 6).value for row in range(9, count + 9)))
                self.assertEqual(2 * count, sum(worksheet.cell(row, 7).value for row in range(9, count + 9)))
                self.assertEqual(list(range(1, count + 1)), [worksheet.cell(row, 1).value for row in range(9, count + 9)])
                self.assertEqual(0, worksheet.page_setup.fitToHeight)
                self.assertEqual(1, worksheet.page_setup.fitToWidth)
                self.assertEqual("$7:$8", worksheet.print_title_rows)
                self.assertEqual("FINISH TIME:____________________________________", worksheet.cell(count + 10, 2).value)
