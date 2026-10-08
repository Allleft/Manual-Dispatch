import io
import json
import os
import unittest
from copy import deepcopy
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from playwright.sync_api import expect

from tests import test_delivery_per_trip_browser_acceptance as browser_fixture
from tests.manual_dispatch_api_test_helpers import create_legacy_combined_delivery_fixture
from tests.test_delivery_per_trip_persistence import DATE


class DeliveryPerTripExportBrowserAcceptanceTest(unittest.TestCase):
    # Reuse the isolated real-app fixture without collecting Stage 3 test methods again.
    fixture = browser_fixture.DeliveryPerTripBrowserAcceptanceTest
    setUpClass = classmethod(fixture.setUpClass.__func__)
    tearDownClass = classmethod(fixture.tearDownClass.__func__)
    setUp = fixture.setUp
    panel = fixture.panel
    selected = fixture.selected
    choose = fixture.choose
    generate = fixture.generate

    def serve(self, route):
        request = route.request
        url = urlsplit(request.url)
        if url.hostname != "per-trip-acceptance.test" or not url.path.startswith("/api/"):
            return self.fixture.serve(self, route)
        payload = request.post_data_json if request.post_data else None
        self.requests.append((request.method, url.path, payload))
        response = self.client.request(request.method, url.path + ("?" + url.query if url.query else ""), json=payload)
        if response.status_code >= 400:
            self.errors.append(f"{request.method} {url.path}: {response.status_code} {response.text}")
        route.fulfill(status=response.status_code, body=response.content, headers=dict(response.headers))

    def download(self, button, filename):
        with self.page.expect_download() as waiting:
            button.click()
        download = waiting.value
        self.assertEqual(filename, download.suggested_filename)
        workbook = load_workbook(io.BytesIO(Path(download.path()).read_bytes()))
        self.addCleanup(workbook.close)
        return workbook

    def history(self):
        self.page.get_by_role("link", name="Saved History", exact=True).click()
        expect(self.page.get_by_role("heading", name="Saved Run Sheet History", exact=True)).to_be_visible()
        self.page.get_by_label("Delivery date", exact=True).fill(DATE)
        self.page.get_by_label("Delivery date", exact=True).dispatch_event("change")
        return self.page.locator(".workspace-history-results .workspace-daily-run-sheet")

    def evidence(self, label, sheets):
        output = os.environ.get("MANUAL_DISPATCH_STAGE4_ACCEPTANCE_OUTPUT")
        if output:
            target = Path(output)
            target.mkdir(parents=True, exist_ok=True)
            self.page.screenshot(path=str(target / f"{label}.png"), full_page=True)
            (target / f"{label}.json").write_text(json.dumps({
                "result": "PASS", "db": str(self.repository.db_path),
                "logbook": str(self.service.logbook.base_dir),
                "run_sheet_ids": [sheet.run_sheet_id for sheet in sheets], "page_errors": self.errors,
            }, indent=2), encoding="utf-8")

    def test_actual_downloads_history_and_independent_closed_open_siblings(self):
        self.choose("trip1", "V001")
        self.assertIn("V001", self.panel("trip2").locator("option").evaluate_all("nodes => nodes.map(node => node.value)"))
        self.choose("trip2", "V001")
        self.choose("trip2", "V002")
        second = self.generate("trip2")
        first = self.generate("trip1")
        self.page.goto("http://per-trip-acceptance.test/#delivery/run-sheet", wait_until="networkidle")
        cards = self.page.locator(".workspace-run-sheet-document-card")
        expect(cards).to_have_count(2)
        self.assertEqual(["John — Trip 1", "John — Trip 2"], cards.locator(".workspace-record-card-top h3").all_text_contents())
        for number, expected in ((1, first), (2, second)):
            card = cards.nth(number - 1)
            expect(card.get_by_text(expected.vehicle_rego_snapshot, exact=False)).to_be_visible()
            card.get_by_role("button", name="Save Run Sheet", exact=True).click()
            expect(card.get_by_role("button", name="Export Excel", exact=True)).to_be_enabled()
            workbook = self.download(card.get_by_role("button", name="Export Excel", exact=True),
                                     f"Delivery_Run_Sheet_{DATE}_John_Trip_{number}.xlsx")
            self.assertEqual([f"John - Trip {number}"], workbook.sheetnames)
            self.assertEqual(f"REGO #: {expected.vehicle_rego_snapshot}", workbook.active["L1"].value)
            self.assertEqual(f"INV-100{number}", workbook.active["D9"].value)
            self.assertIsNone(workbook.active["D10"].value)
        date_workbook = self.download(self.page.get_by_role("button", name="Export Excel File", exact=True),
                                      f"Daily_Run_Sheets_{DATE}.xlsx")
        self.assertEqual(["John - Trip 1", "John - Trip 2"], date_workbook.sheetnames)
        for number, expected in ((1, first), (2, second)):
            worksheet = date_workbook[f"John - Trip {number}"]
            self.assertEqual(f"REGO #: {expected.vehicle_rego_snapshot}", worksheet["L1"].value)
            self.assertEqual(f"INV-100{number}", worksheet["D9"].value)
        sibling = deepcopy(self.repository.get_delivery_run_sheet(second.run_sheet_id))
        cards.nth(0).get_by_role("button", name="Close Run Sheet", exact=True).click()
        modal = self.page.locator(".workspace-modal-backdrop")
        modal.get_by_role("button", name="Mark All Delivered", exact=True).click()
        modal.get_by_role("button", name="Review and Close", exact=True).click()
        expect(modal).to_have_count(0)
        expect(cards.nth(0).locator(".workspace-run-sheet-badges")).to_contain_text("CLOSED")
        expect(cards.nth(1).locator(".workspace-run-sheet-badges")).to_contain_text("OPEN")
        self.assertEqual(sibling, self.repository.get_delivery_run_sheet(second.run_sheet_id))
        papers = self.history()
        try:
            expect(papers).to_have_count(2)
        except AssertionError as error:
            self.fail(f"{error}\nPage: {self.page.locator('body').inner_text()}\nErrors: {self.errors}")
        for number, expected in ((1, first), (2, second)):
            paper = papers.nth(number - 1)
            expect(paper.locator(".workspace-daily-run-sheet-driver")).to_contain_text(f"Trip: Trip {number}")
            expect(paper).to_contain_text(expected.vehicle_rego_snapshot)
            expect(paper).to_contain_text("SAVED")
            expect(paper).to_contain_text("CLOSED" if number == 1 else "OPEN")
        expect(papers.nth(0).locator(".workspace-run-sheet-outcomes")).to_have_count(1)
        expect(papers.nth(1).locator(".workspace-run-sheet-outcomes")).to_have_count(0)
        self.assertEqual([], self.errors)
        self.evidence("per-trip-export-history", [first, second])

    def test_legacy_combined_browser_export_history_and_snapshot_stay_intact(self):
        self.choose("trip1", "V001")
        self.choose("trip2", "V001")
        legacy = create_legacy_combined_delivery_fixture(self.service, DATE, "D001")
        self.page.goto("http://per-trip-acceptance.test/#delivery/run-sheet", wait_until="networkidle")
        cards = self.page.locator(".workspace-run-sheet-document-card")
        expect(cards).to_have_count(1)
        expect(cards.locator(".workspace-record-card-top h3")).to_have_text("John")
        self.assertNotIn("Trip 1", cards.inner_text())
        self.assertNotIn("Trip 2", cards.inner_text())
        cards.get_by_role("button", name="Save Run Sheet", exact=True).click()
        expect(cards.get_by_role("button", name="Export Excel", exact=True)).to_be_enabled()
        saved = deepcopy(self.repository.get_delivery_run_sheet(legacy.run_sheet_id))
        single = self.download(cards.get_by_role("button", name="Export Excel", exact=True),
                               f"Delivery_Run_Sheet_{DATE}_John.xlsx")
        self.assertEqual(["Daily Run Sheet"], single.sheetnames)
        self.assertEqual(["INV-1001", "INV-1002"], [single.active[f"D{row}"].value for row in (9, 10)])
        self.assertIsNone(single.active["F2"].value)
        date_workbook = self.download(self.page.get_by_role("button", name="Export Excel File", exact=True),
                                      f"Daily_Run_Sheets_{DATE}.xlsx")
        self.assertEqual(["John"], date_workbook.sheetnames)
        papers = self.history()
        expect(papers).to_have_count(1)
        expect(papers).to_contain_text("INV-1001")
        expect(papers).to_contain_text("INV-1002")
        self.assertNotIn("Trip 1", papers.inner_text())
        self.assertNotIn("Trip 2", papers.inner_text())
        self.assertEqual(saved, self.repository.get_delivery_run_sheet(legacy.run_sheet_id))
        self.assertEqual([], self.errors)
        self.evidence("legacy-combined-export-history", [saved])
