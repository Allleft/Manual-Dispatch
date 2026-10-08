import importlib
import io
import json
import os
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from dataclasses import replace
from threading import Barrier
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.errors import DeliveryRunSheetLockedError, StateChangedConflictError
from backend.repositories.sqlite_manual_dispatch_repository import SQLiteManualDispatchRepository
from backend.schemas import (
    CloseDeliveryRunSheetRequest, CloseDeliveryRunSheetRowRequest,
    DeliveryWorkspaceAssignOrderRequest, DeliveryWorkspaceUnassignOrderRequest,
    GenerateDeliveryRunSheetRequest, OperatorAccountIdentity,
    RegisterOperatorAccountRequest, SaveGeneratedWorkspaceSnapshotRequest,
)
from backend.services.delivery_run_sheet_excel_export_service import (
    build_delivery_run_sheet_excel, build_delivery_run_sheets_excel,
)
from backend.services.manual_dispatch_service import ManualDispatchService
from tests.manual_dispatch_api_test_helpers import authenticate_test_client
from tests.test_delivery_per_trip_persistence import DATE, snapshot
from tests.test_delivery_per_trip_vehicle_assignment import VehicleFixture


class DeliveryPerTripRunSheetTest(VehicleFixture):
    def setUp(self):
        super().setUp()
        for number in (1, 2):
            order = self.repository.get_order(f"ORD-00{number}")
            self.repository.update_order(replace(
                order, pallet_quantity=number, loose_bags_quantity=number * 3,
                carton_quantity=number * 5,
            ))
        account = self.service.register_operator_account(RegisterOperatorAccountRequest(
            account_name="Trip Test", password="TripTest123!", confirm_password="TripTest123!",
        ))
        self.identity = OperatorAccountIdentity(account.account_id, account.account_name)

    def assign_orders(self, trips=("trip1", "trip2")):
        for number, trip in enumerate(("trip1", "trip2"), 1):
            if trip in trips:
                self.service.assign_delivery_workspace_order(DeliveryWorkspaceAssignOrderRequest(
                    order_id=f"ORD-00{number}", driver_id="D001", trip_no=trip,
                ))

    def generate(self, trip="trip1", dispatch=DATE):
        return self.service.create_generated_delivery_run_sheet(GenerateDeliveryRunSheetRequest(
            delivery_date=DATE, dispatch_date=dispatch, driver_id="D001", trip_no=trip,
        ))

    def save(self, sheet):
        return self.service.save_generated_delivery_run_sheet(
            sheet.run_sheet_id, SaveGeneratedWorkspaceSnapshotRequest(
                saved_by_account_id=self.identity.account_id,
                saved_by_account_name=self.identity.account_name,
            ),
        )

    def close(self, sheet, outcome="DELIVERED"):
        rows = [CloseDeliveryRunSheetRowRequest(
            run_sheet_row_id=row.row_id, outcome=outcome,
            reason_code="TIME_RAN_OUT" if outcome == "RETURN_TO_POOL" else None,
            next_delivery_date="2026-09-16" if outcome == "RETURN_TO_POOL" else None,
        ) for trip in sheet.trips for row in trip.orders]
        return self.service.close_saved_delivery_run_sheet(
            sheet.run_sheet_id, CloseDeliveryRunSheetRequest(rows), self.identity,
        )

    def get(self, sheet):
        return self.repository.get_delivery_run_sheet(sheet.run_sheet_id)

    def test_only_trip1_generates_one_scoped_snapshot(self):
        self.assign_orders(("trip1",))
        sheet = self.generate()
        self.assertEqual("trip1", sheet.trip_no)
        self.assertEqual(["trip1"], [trip.trip_no for trip in sheet.trips])
        self.assertEqual(["ORD-001"], [row.task_id for row in sheet.trips[0].orders])

    def test_only_trip2_generates_one_scoped_snapshot(self):
        self.assign_orders(("trip2",))
        sheet = self.generate("trip2")
        self.assertEqual("trip2", sheet.trip_no)
        self.assertEqual(["trip2"], [trip.trip_no for trip in sheet.trips])
        self.assertEqual(["ORD-002"], [row.task_id for row in sheet.trips[0].orders])

    def test_both_trips_generate_distinct_independent_records(self):
        self.assign_orders()
        first, second = self.generate(), self.generate("trip2")
        self.assertNotEqual(first.run_sheet_id, second.run_sheet_id)
        self.assertEqual(first, self.get(first))
        self.assertEqual({"trip1", "trip2"}, {s.trip_no for s in self.service.list_delivery_run_sheets(delivery_date=DATE)})

    def test_different_trip_vehicles_are_snapshotted_independently(self):
        self.assign_orders()
        self.assign()
        self.assign(trip="trip2", vehicle="V002")
        self.assertEqual(["V001", "V002"], [self.generate(trip).vehicle_id for trip in ("trip1", "trip2")])

    def test_same_vehicle_across_trips_is_valid(self):
        self.assign_orders()
        self.assign()
        self.assign(trip="trip2")
        self.assertEqual(["V001", "V001"], [self.generate(trip).vehicle_id for trip in ("trip1", "trip2")])

    def test_trip1_totals_exclude_trip2(self):
        self.assign_orders()
        sheet = self.generate()
        self.assertEqual((1, 3, 5), (sheet.total_pallets, sheet.total_loose_bags, sheet.total_cartons))

    def test_trip2_totals_exclude_trip1(self):
        self.assign_orders()
        sheet = self.generate("trip2")
        self.assertEqual((2, 6, 10), (sheet.total_pallets, sheet.total_loose_bags, sheet.total_cartons))

    def test_selected_builder_does_not_load_sibling_order(self):
        self.assign_orders()
        original = self.repository.get_order
        def selected_only(order_id):
            self.assertNotEqual("ORD-002", order_id)
            return original(order_id)
        with patch.object(self.repository, "get_order", side_effect=selected_only):
            self.generate()

    def test_duplicate_trip_rejected_across_dispatch_provenance(self):
        self.assign_orders()
        first = self.generate(dispatch="2026-09-01")
        with self.assertRaises(StateChangedConflictError):
            self.generate(dispatch="2026-09-02")
        self.assertEqual(first, self.get(first))
        self.assertEqual("trip2", self.generate("trip2", "2026-09-02").trip_no)

    def test_legacy_combined_blocks_both_trips_without_rewriting(self):
        self.assign_orders()
        legacy = snapshot()
        self.repository.create_delivery_run_sheet(legacy)
        for trip in ("trip1", "trip2"):
            with self.subTest(trip=trip), self.assertRaisesRegex(StateChangedConflictError, "legacy combined"):
                self.generate(trip)
        self.assertEqual(legacy, self.get(legacy))

    def test_generation_requires_trip_without_default(self):
        with self.assertRaises(TypeError):
            GenerateDeliveryRunSheetRequest(delivery_date=DATE, driver_id="D001")
        self.assign_orders()
        for trip in (None, "", "trip3", "arbitrary"):
            with self.subTest(trip=trip), self.assertRaises(ValueError):
                self.generate(trip)
        self.assertEqual([], self.service.list_delivery_run_sheets(delivery_date=DATE))

    def test_empty_selected_trip_rejected_even_with_sibling_orders(self):
        self.assign_orders(("trip2",))
        with self.assertRaisesRegex(ValueError, "assigned Delivery Order"):
            self.generate()

    def test_vehicle_remains_optional_and_null_day_selection_is_not_projected(self):
        self.assign_orders()
        self.legacy()
        first = self.generate()
        self.assertIsNone(first.vehicle_id)
        self.assertIsNone(first.vehicle_rego_snapshot)

    def test_every_new_snapshot_row_matches_header_trip(self):
        self.assign_orders()
        for trip in ("trip1", "trip2"):
            sheet = self.generate(trip)
            self.assertEqual(1, len(sheet.trips))
            self.assertTrue(all(row.trip_no == sheet.trip_no for row in sheet.trips[0].orders))

    def test_repository_rejects_mixed_trip_new_snapshot(self):
        legacy = snapshot()
        with self.assertRaises(ValueError):
            self.repository.create_delivery_run_sheet(replace(legacy, trip_no="trip1"))
        single = snapshot(trip_no="trip1")
        single.trips[0].orders[0].trip_no = "trip2"
        with self.assertRaises(ValueError):
            self.repository.create_delivery_run_sheet(single)

    def test_save_trip1_leaves_ungenerated_trip2_editable(self):
        self.assign_orders()
        self.save(self.generate())
        self.assign(trip="trip2", vehicle="V002")
        self.service.unassign_delivery_workspace_order(DeliveryWorkspaceUnassignOrderRequest(order_id="ORD-002"))
        self.assertIsNone(self.repository.find_assignment_for_task("ORDER", "ORD-002"))

    def test_save_trip1_does_not_change_trip2_sheet(self):
        self.assign_orders()
        first, second = self.generate(), self.generate("trip2")
        self.save(first)
        self.assertEqual(second, self.get(second))

    def test_cancel_trip1_preserves_trip2_sheet_vehicle_orders(self):
        self.assign_orders()
        self.assign()
        self.assign(trip="trip2", vehicle="V002")
        first, second = self.generate(), self.generate("trip2")
        order = deepcopy(self.repository.get_order("ORD-002"))
        assignment = deepcopy(self.repository.find_assignment_for_task("ORDER", "ORD-002"))
        self.service.cancel_generated_delivery_run_sheet(first.run_sheet_id)
        self.assertIsNone(self.get(first))
        self.assertEqual(second, self.get(second))
        self.assertEqual(order, self.repository.get_order("ORD-002"))
        self.assertEqual(assignment, self.repository.find_assignment_for_task("ORDER", "ORD-002"))
        self.assertEqual("V002", self.selected(trip="trip2"))
        self.assign(vehicle="V003")

    def test_close_trip1_leaves_saved_open_trip2_and_vehicle_unchanged(self):
        self.assign_orders()
        self.assign()
        self.assign(trip="trip2", vehicle="V002")
        first, second = self.save(self.generate()), self.save(self.generate("trip2"))
        sibling = deepcopy(self.repository.get_order("ORD-002"))
        self.assertEqual("CLOSED", self.close(first).execution_status)
        self.assertEqual(second, self.get(second))
        self.assertEqual(sibling, self.repository.get_order("ORD-002"))
        self.assertEqual("V002", self.selected(trip="trip2"))

    def test_close_trip2_leaves_trip1_sheet_order_vehicle_unchanged(self):
        self.assign_orders()
        self.assign()
        first, second = self.save(self.generate()), self.save(self.generate("trip2"))
        sibling = deepcopy(self.repository.get_order("ORD-001"))
        self.close(second)
        self.assertEqual(first, self.get(first))
        self.assertEqual(sibling, self.repository.get_order("ORD-001"))
        self.assertEqual("V001", self.selected())

    def test_close_trip1_leaves_ungenerated_trip2_editable(self):
        self.assign_orders()
        self.close(self.save(self.generate()))
        self.assign(trip="trip2", vehicle="V002")
        self.service.unassign_delivery_workspace_order(DeliveryWorkspaceUnassignOrderRequest(order_id="ORD-002"))
        self.assertEqual("ACTIVE", self.repository.get_order("ORD-002").status)

    def test_return_to_pool_updates_only_selected_order(self):
        self.assign_orders()
        first, second = self.save(self.generate()), self.save(self.generate("trip2"))
        sibling_order = deepcopy(self.repository.get_order("ORD-002"))
        sibling_assignment = deepcopy(self.repository.find_assignment_for_task("ORDER", "ORD-002"))
        self.close(first, "RETURN_TO_POOL")
        returned = self.repository.get_order("ORD-001")
        self.assertEqual(("ACTIVE", "2026-09-16"), (returned.status, returned.delivery_date))
        self.assertIsNone(self.repository.find_assignment_for_task("ORDER", "ORD-001"))
        self.assertEqual(sibling_order, self.repository.get_order("ORD-002"))
        self.assertEqual(sibling_assignment, self.repository.find_assignment_for_task("ORDER", "ORD-002"))
        self.assertEqual(second, self.get(second))

    def test_generated_lock_protects_source_and_target_but_not_sibling_edits(self):
        self.assign_orders()
        self.generate()
        for order_id, trip in (("ORD-001", "trip2"), ("ORD-002", "trip1")):
            with self.subTest(order_id=order_id), self.assertRaises(DeliveryRunSheetLockedError):
                self.service.assign_delivery_workspace_order(DeliveryWorkspaceAssignOrderRequest(
                    order_id=order_id, driver_id="D001", trip_no=trip,
                ))
        self.service.unassign_delivery_workspace_order(DeliveryWorkspaceUnassignOrderRequest(order_id="ORD-002"))
        self.assign(trip="trip2", vehicle="V002")
        with self.assertRaises(DeliveryRunSheetLockedError):
            self.clear()

    def test_legacy_combined_lock_blocks_source_and_both_targets(self):
        self.assign_orders()
        legacy = snapshot()
        self.repository.create_delivery_run_sheet(legacy)
        for number, trip in enumerate(("trip1", "trip2"), 1):
            with self.subTest(trip=trip), self.assertRaises(DeliveryRunSheetLockedError):
                self.service.unassign_delivery_workspace_order(DeliveryWorkspaceUnassignOrderRequest(order_id=f"ORD-00{number}"))
            with self.subTest(trip=trip), self.assertRaises(DeliveryRunSheetLockedError):
                self.assign(trip=trip)

    def test_closeout_rejects_new_header_row_mismatch_before_order_writes(self):
        self.assign_orders()
        sheet = self.save(self.generate())
        malformed = deepcopy(sheet)
        malformed.trips[0].orders[0].trip_no = "trip2"
        before = deepcopy(self.repository.get_order("ORD-001"))
        with patch.object(self.repository, "get_delivery_run_sheet", return_value=malformed):
            with self.assertRaisesRegex(StateChangedConflictError, "inconsistent trip snapshot"):
                self.close(malformed)
        self.assertEqual(before, self.repository.get_order("ORD-001"))
        self.assertEqual(sheet, self.get(sheet))

    def test_legacy_combined_save_close_and_export_preserve_frozen_snapshot(self):
        self.assign_orders()
        legacy = snapshot()
        self.repository.create_delivery_run_sheet(legacy)
        saved = self.save(legacy)
        workbook = load_workbook(io.BytesIO(build_delivery_run_sheet_excel(saved)))
        self.assertEqual(1, len(workbook.worksheets))
        workbook.close()
        closed = self.close(saved)
        self.assertIsNone(closed.trip_no)
        self.assertEqual(legacy.trips, closed.trips)
        self.assertEqual(legacy.vehicle_rego_snapshot, closed.vehicle_rego_snapshot)
        self.assertEqual("CLOSED", closed.execution_status)

    def test_date_export_keeps_both_trips_using_existing_unique_sheet_names(self):
        self.assign_orders()
        first, second = self.generate(), self.generate("trip2")
        sheets = self.service.list_delivery_run_sheets(delivery_date=DATE)
        workbook = load_workbook(io.BytesIO(build_delivery_run_sheets_excel(sheets, DATE)))
        self.assertEqual(2, len(workbook.worksheets))
        self.assertEqual(2, len(set(workbook.sheetnames)))
        contents = [str(list(sheet.values)) for sheet in workbook.worksheets]
        for sheet in (first, second):
            invoice = sheet.trips[0].orders[0].invoice_number_snapshot
            self.assertEqual(1, sum(invoice in content for content in contents))
        workbook.close()

    def test_unavailable_driver_existing_work_generates_saves_and_closes(self):
        self.assign_orders()
        self.repository.update_driver(replace(self.repository.get_driver("D001"), is_available=False))
        first = self.save(self.generate())
        self.close(first)
        self.assertEqual("trip2", self.generate("trip2").trip_no)

    def test_lifecycle_logbook_events_include_trip_identity(self):
        self.assign_orders()
        first, second = self.save(self.generate()), self.generate("trip2")
        self.close(first)
        self.service.cancel_generated_delivery_run_sheet(second.run_sheet_id)
        entries = [json.loads(line) for path in self.logbook.base_dir.glob("*.txt")
                   for line in path.read_text(encoding="utf-8").splitlines()]
        actions = {"DELIVERY_RUN_SHEET_GENERATED", "DELIVERY_RUN_SHEET_SAVED", "DELIVERY_RUN_SHEET_CLOSED", "DELIVERY_RUN_SHEET_CANCELLED"}
        events = [event for event in entries if event["action"] in actions and event["result"] == "SUCCESS"]
        self.assertEqual(actions, {event["action"] for event in events})
        for event in events:
            self.assertEqual("trip1" if event["run_sheet_id"] == first.run_sheet_id else "trip2", event["metadata"]["trip_no"])


class DeliveryPerTripRunSheetSQLiteTest(DeliveryPerTripRunSheetTest):
    def build_repository(self):
        with patch.dict(os.environ, {"MANUAL_DISPATCH_SEED_DEMO_DATA": "true"}):
            return SQLiteManualDispatchRepository(self.root / "db.sqlite3")

    def test_concurrent_same_trip_generation_has_one_winner(self):
        self.assign_orders()
        self._race(("trip1", "trip1"), expected_winners=1)

    def test_concurrent_sibling_trip_generation_has_two_winners(self):
        self.assign_orders()
        self._race(("trip1", "trip2"), expected_winners=2)

    def _race(self, trips, expected_winners):
        barrier = Barrier(2)
        def generate(trip):
            repository = SQLiteManualDispatchRepository(self.repository.db_path)
            service = ManualDispatchService(repository, self.logbook)
            barrier.wait(timeout=10)
            try:
                return service.create_generated_delivery_run_sheet(GenerateDeliveryRunSheetRequest(
                    trip_no=trip, delivery_date=DATE, driver_id="D001",
                ))
            except StateChangedConflictError:
                return None
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(generate, trips))
        self.assertEqual(expected_winners, sum(result is not None for result in results))
        self.assertEqual(expected_winners, len(self.service.list_delivery_run_sheets(delivery_date=DATE)))


class DeliveryPerTripRunSheetApiTest(VehicleFixture):
    build_repository = DeliveryPerTripRunSheetSQLiteTest.build_repository

    def setUp(self):
        super().setUp()
        api = importlib.import_module("backend.api.manual_dispatch")
        previous = api.service
        api.service = self.service
        self.addCleanup(setattr, api, "service", previous)
        app = FastAPI()
        app.include_router(api.router)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        authenticate_test_client(self.client, self.service)
        for number, trip in enumerate(("trip1", "trip2"), 1):
            self.service.assign_delivery_workspace_order(DeliveryWorkspaceAssignOrderRequest(
                order_id=f"ORD-00{number}", driver_id="D001", trip_no=trip,
            ))

    def test_api_rejects_missing_null_empty_and_unsupported_trip(self):
        payload = {"delivery_date": DATE, "driver_id": "D001"}
        path = "/api/manual-dispatch/delivery/run-sheets/generated"
        self.assertEqual(422, self.client.post(path, json=payload).status_code)
        self.assertEqual(422, self.client.post(path, json={**payload, "trip_no": None}).status_code)
        for trip in ("", "trip3", "arbitrary"):
            with self.subTest(trip=trip):
                self.assertEqual(400, self.client.post(path, json={**payload, "trip_no": trip}).status_code)

    def test_api_returns_both_headers_and_history_without_driver_date_dedupe(self):
        ids = []
        for trip in ("trip1", "trip2"):
            response = self.client.post("/api/manual-dispatch/delivery/run-sheets/generated", json={
                "delivery_date": DATE, "driver_id": "D001", "trip_no": trip,
            })
            self.assertEqual(200, response.status_code, response.text)
            sheet = response.json()
            self.assertEqual(trip, sheet["trip_no"])
            self.assertEqual([trip], [item["trip_no"] for item in sheet["trips"]])
            ids.append(sheet["run_sheet_id"])
            saved = self.client.post(f"/api/manual-dispatch/delivery/run-sheets/{ids[-1]}/save", json={})
            self.assertEqual(200, saved.status_code, saved.text)
        listed = self.client.get("/api/manual-dispatch/delivery/run-sheets", params={"delivery_date": DATE, "status": "SAVED"})
        self.assertEqual(200, listed.status_code)
        self.assertEqual(set(ids), {item["run_sheet_id"] for item in listed.json()})
        self.assertEqual({"trip1", "trip2"}, {item["trip_no"] for item in listed.json()})
