import importlib
import io
import os
import tempfile
import unittest
from copy import deepcopy
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.repositories.in_memory_manual_dispatch_repository import InMemoryManualDispatchRepository
from backend.repositories.sqlite_manual_dispatch_repository import SQLiteManualDispatchRepository
from backend.services.delivery_run_sheet_excel_export_service import (
    build_delivery_run_sheet_excel, build_delivery_run_sheets_excel,
)
from backend.services.manual_dispatch.logbook_file_service import LogbookFileService
from backend.services.manual_dispatch_service import ManualDispatchService
from tests.manual_dispatch_api_test_helpers import authenticate_test_client
from tests.test_delivery_per_trip_persistence import DATE, snapshot


def sheet(driver="D001", name="John", trip="trip1", identity=None, **changes):
    return replace(snapshot(identity or f"{driver}-{trip}", trip_no=trip, driver=driver),
                   driver_name_snapshot=name, **changes)


class DeliveryPerTripExcelPresentationTest(unittest.TestCase):
    def workbook(self, sheets, single=False):
        content = (build_delivery_run_sheet_excel(sheets[0]) if single else
                   build_delivery_run_sheets_excel(sheets, DATE))
        workbook = load_workbook(io.BytesIO(content))
        self.addCleanup(workbook.close)
        return workbook

    def test_only_actual_trip1_creates_one_worksheet(self):
        workbook = self.workbook([sheet()])
        self.assertEqual(["John - Trip 1"], workbook.sheetnames)

    def test_trip1_trip2_names_and_order_ignore_input_and_generated_time(self):
        first = sheet(generated_at="2030-01-01T01:00:00Z")
        second = sheet(trip="trip2", generated_at="2020-01-01T01:00:00Z")
        self.assertEqual(["John - Trip 1", "John - Trip 2"], self.workbook([second, first]).sheetnames)

    def test_same_display_name_drivers_get_stable_unique_names(self):
        rows = [sheet("D002"), sheet("D001", trip="trip2"), sheet()]
        workbook = self.workbook(rows)
        self.assertEqual(["John - Trip 1", "John - Trip 2", "John (2) - Trip 1"], workbook.sheetnames)
        self.assertEqual(workbook.sheetnames, self.workbook(list(reversed(rows))).sheetnames)

    def test_case_insensitive_name_collisions_remain_unique(self):
        rows = [sheet("D002", "john"), sheet("D001", "John")]
        self.assertEqual(["John - Trip 1", "john (2) - Trip 1"], self.workbook(rows).sheetnames)

    def test_long_names_reserve_trip_suffix_and_duplicate_number(self):
        rows = [sheet("D002", "Very long shared driver name " * 3),
                sheet("D001", "Very long shared driver name " * 3),
                sheet("D001", "Very long shared driver name " * 3, "trip2")]
        names = self.workbook(rows).sheetnames
        self.assertTrue(all(len(name) <= 31 for name in names))
        self.assertEqual(3, len({name.casefold() for name in names}))
        self.assertTrue(names[0].endswith(" - Trip 1"))
        self.assertTrue(names[1].endswith(" - Trip 2"))
        self.assertTrue(names[2].endswith(" (2) - Trip 1"))

    def test_invalid_excel_characters_are_sanitized(self):
        name = "'John \\/ * ? : [North]'"
        workbook = self.workbook([sheet(name=name)])
        self.assertEqual(["John North - Trip 1"], workbook.sheetnames)

    def test_sanitized_names_collide_deterministically(self):
        rows = [sheet("D002", "John:North"), sheet("D001", "John/North")]
        workbook = self.workbook(rows)
        self.assertEqual(["John North - Trip 1", "John North (2) - Trip 1"], workbook.sheetnames)
        self.assertEqual(workbook.sheetnames, self.workbook(list(reversed(rows))).sheetnames)

    def test_legacy_and_per_trip_name_collision_keeps_workbook_valid(self):
        rows = [sheet("D002", "John - Trip 1", None), sheet()]
        workbook = self.workbook(rows)
        self.assertEqual(["John - Trip 1", "John - Trip 1 (2)"], workbook.sheetnames)
        self.assertEqual(workbook.sheetnames, self.workbook(list(reversed(rows))).sheetnames)

    def test_per_trip_collision_keeps_suffix_when_legacy_name_sorts_first(self):
        rows = [sheet("D001", "A - Trip 1", None), sheet("D002", "A\\", "trip1")]
        names = self.workbook(rows).sheetnames
        self.assertEqual(2, len({name.casefold() for name in names}))
        self.assertTrue(all(len(name) <= 31 for name in names))
        self.assertTrue(names[1].endswith(" - Trip 1"))
        self.assertEqual(["A - Trip 1", "A (2) - Trip 1"], names)

    def test_multiple_drivers_sort_name_then_id_then_trip(self):
        rows = [sheet("D002", "John", "trip2"), sheet("D003", "gavin"),
                sheet("D001", "JOHN", "trip2"), sheet("D002", "John"), sheet("D001", "JOHN")]
        self.assertEqual(["gavin - Trip 1", "JOHN - Trip 1", "JOHN - Trip 2",
                          "John (2) - Trip 1", "John (2) - Trip 2"], self.workbook(rows).sheetnames)

    def test_trip_content_and_different_rego_snapshots_never_mix(self):
        first = sheet(vehicle_rego_snapshot="TRUCK A")
        second = sheet(trip="trip2", vehicle_rego_snapshot="TRUCK B")
        second.trips[0].orders[0].invoice_number_snapshot = "TRIP2-ONLY"
        original = deepcopy([first, second])
        workbook = self.workbook([second, first])
        for number, expected in ((1, first), (2, second)):
            worksheet = workbook[f"John - Trip {number}"]
            self.assertEqual(f"REGO #: TRUCK {'A' if number == 1 else 'B'}", worksheet["L1"].value)
            self.assertEqual(expected.trips[0].orders[0].invoice_number_snapshot, worksheet["D9"].value)
            self.assertIsNone(worksheet["D10"].value)
            self.assertEqual(f"TRIP: TRIP {number}", worksheet["F2"].value)
        self.assertEqual(original, [first, second])

    def test_same_vehicle_is_printed_on_both_actual_trips(self):
        workbook = self.workbook([sheet(), sheet(trip="trip2")])
        self.assertEqual(["REGO #: ORIGINAL"] * 2, [row["L1"].value for row in workbook.worksheets])

    def test_single_exports_have_one_named_trip_sheet_and_header(self):
        for trip, number in (("trip1", 1), ("trip2", 2)):
            with self.subTest(trip=trip):
                workbook = self.workbook([sheet(trip=trip)], single=True)
                self.assertEqual([f"John - Trip {number}"], workbook.sheetnames)
                self.assertEqual(f"TRIP: TRIP {number}", workbook.active["F2"].value)
                self.assertEqual("DAILY RUN SHEET", workbook.active["A1"].value)
                self.assertEqual("landscape", workbook.active.page_setup.orientation)

    def test_legacy_combined_keeps_single_sheet_names_and_paper_layout(self):
        legacy = sheet(trip=None)
        before = deepcopy(legacy)
        single = self.workbook([legacy], single=True)
        date = self.workbook([legacy])
        self.assertEqual(["Daily Run Sheet"], single.sheetnames)
        self.assertEqual(["John"], date.sheetnames)
        for workbook in (single, date):
            worksheet = workbook.active
            self.assertIsNone(worksheet["F2"].value)
            self.assertEqual(["INV-1", "INV-2"], [worksheet[f"D{row}"].value for row in (9, 10)])
            self.assertEqual("DRIVER: John", worksheet["F1"].value)
            self.assertEqual("$7:$8", worksheet.print_title_rows)
        self.assertEqual(before, legacy)


class DeliveryPerTripListingPresentationTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="delivery-stage4-list-")
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        with patch.dict(os.environ, {"MANUAL_DISPATCH_SEED_DEMO_DATA": "true"}):
            self.repositories = [InMemoryManualDispatchRepository(), SQLiteManualDispatchRepository(root / "isolated.sqlite3")]
        self.services = [ManualDispatchService(repo, LogbookFileService(root / f"logbook-{number}"))
                         for number, repo in enumerate(self.repositories)]

    def test_sqlite_and_memory_list_and_date_export_have_equivalent_order(self):
        rows = [sheet("D002", "John", "trip2"), sheet("D003", "gavin"),
                sheet("D001", "JOHN", "trip2"), sheet("D002", "John"), sheet("D001", "JOHN")]
        expected = ["D003-trip1", "D001-trip1", "D001-trip2", "D002-trip1", "D002-trip2"]
        for repo, service in zip(self.repositories, self.services):
            with self.subTest(repository=type(repo).__name__):
                for number, row in enumerate(reversed(rows)):
                    repo.create_delivery_run_sheet(replace(row, generated_at=f"2026-09-14T08:00:0{number}Z"))
                self.assertEqual(expected, [row.run_sheet_id for row in service.list_delivery_run_sheets(delivery_date=DATE)])
                self.assertEqual(expected, [row.run_sheet_id for row in service.list_delivery_run_sheets_for_date_export(DATE)])

    def test_history_keeps_newest_date_first_and_closed_open_siblings(self):
        rows = [sheet("D001", "John", "trip2", status="SAVED"),
                sheet("D001", "John", execution_status="CLOSED", status="SAVED"),
                sheet("D002", "Tony", None, status="SAVED", delivery_date="2026-09-16")]
        for repo, service in zip(self.repositories, self.services):
            with self.subTest(repository=type(repo).__name__):
                for row in rows:
                    repo.create_delivery_run_sheet(row)
                result = service.list_delivery_run_sheets(status="SAVED")
                self.assertEqual(["D002-None", "D001-trip1", "D001-trip2"], [row.run_sheet_id for row in result])
                self.assertEqual([None, "trip1", "trip2"], [row.trip_no for row in result])
                self.assertEqual(["OPEN", "CLOSED", "OPEN"], [row.execution_status for row in result])


class DeliveryPerTripExportFilenameTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="delivery-stage4-filename-")
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        self.repository = InMemoryManualDispatchRepository()
        self.service = ManualDispatchService(self.repository, LogbookFileService(root / "logbook"))
        api = importlib.import_module("backend.api.manual_dispatch")
        previous = api.service
        api.service = self.service
        self.addCleanup(setattr, api, "service", previous)
        app = FastAPI()
        app.include_router(api.router)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        authenticate_test_client(self.client, self.service)

    def export(self, row):
        self.repository.create_delivery_run_sheet(row)
        response = self.client.get(f"/api/manual-dispatch/delivery/run-sheets/{row.run_sheet_id}/export-excel")
        self.assertEqual(200, response.status_code, response.text[:200])
        return response

    def test_single_trip_filenames_and_logbook_keep_date_driver_trip(self):
        for trip, number in (("trip1", 1), ("trip2", 2)):
            with self.subTest(trip=trip):
                response = self.export(sheet(trip=trip, status="SAVED"))
                self.assertEqual(f'attachment; filename="Delivery_Run_Sheet_{DATE}_John_Trip_{number}.xlsx"',
                                 response.headers["content-disposition"])
        import json
        events = [json.loads(line) for path in self.service.logbook.base_dir.glob("*.txt")
                  for line in path.read_text(encoding="utf-8").splitlines()]
        exports = [event for event in events if event["action"] == "DELIVERY_RUN_SHEET_EXPORTED"]
        self.assertEqual({"D001-trip1", "D001-trip2"}, {event["run_sheet_id"] for event in exports})
        for event in exports:
            first = event["run_sheet_id"].endswith("trip1")
            self.assertEqual(int(first), event["metadata"]["trip1_count"])
            self.assertEqual(int(not first), event["metadata"]["trip2_count"])
        self.assertEqual({f"Delivery_Run_Sheet_{DATE}_John_Trip_{number}.xlsx" for number in (1, 2)},
                         {event["metadata"]["filename"] for event in exports})

    def test_filename_uses_existing_filesystem_sanitizer(self):
        response = self.export(sheet(name='John/\\:*?"<>| G', status="SAVED"))
        self.assertEqual(f'attachment; filename="Delivery_Run_Sheet_{DATE}_John_G_Trip_1.xlsx"',
                         response.headers["content-disposition"])

    def test_legacy_filename_is_unchanged(self):
        response = self.export(sheet(trip=None, status="SAVED"))
        self.assertEqual(f'attachment; filename="Delivery_Run_Sheet_{DATE}_John.xlsx"',
                         response.headers["content-disposition"])

    def test_date_filename_remains_date_workbook(self):
        for trip in ("trip2", "trip1"):
            self.repository.create_delivery_run_sheet(sheet(trip=trip, status="SAVED"))
        response = self.client.get("/api/manual-dispatch/delivery/run-sheets/export-excel", params={"delivery_date": DATE})
        self.assertEqual(200, response.status_code)
        self.assertEqual(f'attachment; filename="Daily_Run_Sheets_{DATE}.xlsx"', response.headers["content-disposition"])
        workbook = load_workbook(io.BytesIO(response.content))
        self.addCleanup(workbook.close)
        self.assertEqual(["John - Trip 1", "John - Trip 2"], workbook.sheetnames)

    def test_export_routes_stay_authenticated_and_invalid_id_is_rejected(self):
        self.client.cookies.clear()
        paths = [f"/api/manual-dispatch/delivery/run-sheets/missing/export-excel",
                 f"/api/manual-dispatch/delivery/run-sheets/export-excel?delivery_date={DATE}"]
        for path in paths:
            with self.subTest(path=path):
                self.assertEqual(401, self.client.get(path).status_code)
        authenticate_test_client(self.client, self.service)
        self.assertEqual(404, self.client.get(paths[0]).status_code)
